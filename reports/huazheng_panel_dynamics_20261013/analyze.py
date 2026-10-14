"""Quality checks and rating dynamics for a third-party Huazheng ESG rating panel (2009-2022).

Input: the netdisk copy in data/raw/huazheng_panel_netdisk_20261013/ (not committed). Output: aggregate
statistics only (summary.json). Grades are mapped C=1 ... AAA=9; 0 and blank cells are treated as missing.
"""
import collections
import hashlib
import json
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data/raw/huazheng_panel_netdisk_20261013/huazheng_2009_2022.xlsx"
GRADES = ["C", "CC", "CCC", "B", "BB", "BBB", "A", "AA", "AAA"]
MAP = {g: i + 1 for i, g in enumerate(GRADES)}


def main():
    book = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    quarterly = list(book["季度数据"].iter_rows(values_only=True))
    year_end = list(book["年末数据"].iter_rows(values_only=True))
    numeric = list(book["季度平均"].iter_rows(values_only=True))
    header = [str(h).strip()[:10] for h in quarterly[0][2:]]
    X = np.array([[MAP.get(v, np.nan) for v in row[2:]] for row in quarterly[1:]], dtype=float)

    qi = {h: j + 2 for j, h in enumerate(header)}
    yi = {h: j for j, h in enumerate(year_end[0])}
    year_mismatch = year_total = 0
    for rq, ry in zip(quarterly[1:], year_end[1:]):
        assert rq[0] == ry[0]
        for year in year_end[0][2:]:
            gq, gy = rq[qi[f"{year}-12-31"]], ry[yi[year]]
            if gq in MAP and gy in MAP:
                year_total += 1
                year_mismatch += gq != gy
    ni = {str(h).strip()[:10]: j for j, h in enumerate(numeric[0])}
    numeric_mismatch = numeric_total = altered_names = 0
    for rq, rn in zip(quarterly[1:], numeric[1:]):
        altered_names += rq[1] != rn[1]
        for h, j in qi.items():
            if h in ni and rq[j] in MAP:
                numeric_total += 1
                value = rn[ni[h]]
                numeric_mismatch += not isinstance(value, (int, float)) or abs(value - MAP[rq[j]]) > 1e-9

    change = X[:, 1:] - X[:, :-1]
    ok = np.isfinite(change)
    d = change[ok]
    level = X[:, :-1][ok]
    T = np.zeros((9, 9))
    for i, j in zip(level.astype(int), X[:, 1:][ok].astype(int)):
        T[i - 1, j - 1] += 1
    lag4 = np.isfinite(X[:, 4:]) & np.isfinite(X[:, :-4])
    years = {}
    for year in range(2009, 2023):
        cols = [j for j, h in enumerate(header) if h.startswith(str(year))]
        years[year] = {"mean_grade": float(np.nanmean(X[:, cols])), "rated_firms_q4": int(np.isfinite(X[:, cols[-1]]).sum())}
    values = collections.Counter(v for row in quarterly[1:] for v in row[2:])
    summary = {
        "source": {"file": SOURCE.relative_to(ROOT).as_posix(), "sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
                   "provenance": "Netdisk copy labelled 'Wind database' via the 数据皮皮侠 public account; not verified against Wind or CSMAR."},
        "coverage": {"firms": len(quarterly) - 1, "quarters": len(header), "first": header[0], "last": header[-1],
                     "cell_values": {str(k): v for k, v in values.most_common()}},
        "checks": {"year_end_vs_q4_mismatches": [year_mismatch, year_total],
                   "numeric_sheet_vs_grade_mapping_mismatches": [numeric_mismatch, numeric_total],
                   "numeric_sheet_altered_firm_names": altered_names},
        "dynamics": {"quarter_pairs": int(d.size), "p_change": float(np.mean(d != 0)), "p_up": float(np.mean(d > 0)),
                     "p_down": float(np.mean(d < 0)), "mean_change_notches": float(d.mean()), "sd_change_notches": float(d.std()),
                     "abs_change_counts": {int(k): int(v) for k, v in sorted(collections.Counter(np.abs(d).astype(int)).items())},
                     "corr_lag1q": float(np.corrcoef(X[:, 1:][ok], X[:, :-1][ok])[0, 1]),
                     "corr_lag4q": float(np.corrcoef(X[:, 4:][lag4], X[:, :-4][lag4])[0, 1]),
                     "expected_change_by_grade": {GRADES[k - 1]: float(d[level == k].mean()) for k in range(1, 10) if (level == k).sum() > 100},
                     "transition_counts": T.astype(int).tolist(), "grade_order": GRADES},
        "pooled_mean_grade": float(np.nanmean(X)),
        "yearly": years,
    }
    (Path(__file__).parent / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: summary[k] for k in ("checks", "pooled_mean_grade")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
